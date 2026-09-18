"""Parser das planilhas oficiais da Caixa.

Cada modalidade publica o historico com cabecalhos proprios (Bola1..Bola6, Bola1..Bola15,
Coluna1..Coluna7, Mes da Sorte). A estrategia por modalidade fica em um mapa, e os
cabecalhos passam por normalizacao (strip, minusculo, sem acento, sem separador) antes de
serem resolvidos, porque a Caixa alterna entre "Data Sorteio" e "Data do Sorteio".
"""

from __future__ import annotations

import csv
import hashlib
import html
import io
import re
import unicodedata
from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook

from app.core.errors import ImportValidationError, ValidationError
from app.domain.entities import Draw
from app.domain.enums import Modality
from app.domain.rules import ModalityRule, rule_for

CONTEST_ALIASES = ("concurso", "nrconcurso", "numeroconcurso", "sorteio")
DATE_ALIASES = ("datasorteio", "datadosorteio", "data", "datadoconcurso", "dataapuracao")
MONTH_ALIASES = ("mesdasorte", "mesdesorte", "mes", "messorte")

MONTH_NAMES = {
    "janeiro": 1, "fevereiro": 2, "marco": 3, "abril": 4, "maio": 5, "junho": 6,
    "julho": 7, "agosto": 8, "setembro": 9, "outubro": 10, "novembro": 11, "dezembro": 12,
}


@dataclass(frozen=True, slots=True)
class RowError:
    row: int
    field: str
    message: str
    value: str | None = None

    def as_dict(self) -> dict[str, Any]:
        return {"row": self.row, "field": self.field, "message": self.message, "value": self.value}


@dataclass(slots=True)
class ParseResult:
    draws: list[Draw] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)
    file_hash: str = ""

    @property
    def is_valid(self) -> bool:
        return not self.errors

    @property
    def first_contest(self) -> int | None:
        return min((draw.contest_no for draw in self.draws), default=None)

    @property
    def last_contest(self) -> int | None:
        return max((draw.contest_no for draw in self.draws), default=None)

    @property
    def first_drawn_at(self) -> date | None:
        return min((draw.drawn_at for draw in self.draws), default=None)

    @property
    def last_drawn_at(self) -> date | None:
        return max((draw.drawn_at for draw in self.draws), default=None)


def normalize_header(value: object) -> str:
    # A planilha do Dia de Sorte traz o mes com entidade HTML, como Mar&ccedil;o, entao o
    # texto passa por unescape antes de perder o acento.
    text = html.unescape(str(value or "")).strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]", "", text)


def file_hash(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


@dataclass(frozen=True, slots=True)
class ModalityStrategy:
    """Como ler uma linha da planilha de uma modalidade."""

    value_prefixes: tuple[str, ...]
    value_count: int
    extras_reader: (
        Callable[[dict[str, Any], ModalityRule, Sequence[int]], dict[str, Any]] | None
    ) = None
    allow_repeated_values: bool = False


def _read_month(
    row: dict[str, Any], rule: ModalityRule, numbers: Sequence[int]
) -> dict[str, Any]:
    raw = _first_present(row, MONTH_ALIASES)
    if raw is None or str(raw).strip() == "":
        raise ValidationError("Campo 'Mes da Sorte' ausente.")
    text = normalize_header(raw)
    month = int(text) if text.isdigit() else MONTH_NAMES.get(text, 0)
    if not 1 <= month <= 12:
        raise ValidationError(f"Mes da Sorte invalido: {raw}.")
    return {"month": month}


def _read_columns(
    row: dict[str, Any], rule: ModalityRule, numbers: Sequence[int]
) -> dict[str, Any]:
    """O Super Sete guarda os digitos tambem em extras, na ordem das colunas."""
    return {"columns": list(numbers)}


STRATEGIES: dict[Modality, ModalityStrategy] = {
    Modality.MEGA_SENA: ModalityStrategy(value_prefixes=("bola", "dezena", "n"), value_count=6),
    Modality.LOTOFACIL: ModalityStrategy(value_prefixes=("bola", "dezena", "n"), value_count=15),
    Modality.QUINA: ModalityStrategy(value_prefixes=("bola", "dezena", "n"), value_count=5),
    Modality.DIA_DE_SORTE: ModalityStrategy(
        value_prefixes=("bola", "dezena", "n"), value_count=7, extras_reader=_read_month
    ),
    Modality.SUPER_SETE: ModalityStrategy(
        value_prefixes=("coluna", "col"),
        value_count=7,
        extras_reader=_read_columns,
        allow_repeated_values=True,
    ),
}


def _first_present(row: dict[str, Any], aliases: Iterable[str]) -> Any:
    for alias in aliases:
        if alias in row and row[alias] not in (None, ""):
            return row[alias]
    return None


def _collect_values(row: dict[str, Any], prefixes: Sequence[str], count: int) -> list[int]:
    for prefix in prefixes:
        keys = [f"{prefix}{index}" for index in range(1, count + 1)]
        if all(key in row for key in keys):
            return [_as_int(row[key]) for key in keys]
    raise ValidationError(
        "Nao foi possivel localizar as colunas de dezenas. "
        f"Esperado {count} colunas com prefixo {prefixes[0]}."
    )


def _as_int(value: object) -> int:
    if value is None or str(value).strip() == "":
        raise ValidationError("Dezena vazia.")
    text = str(value).strip()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if not re.fullmatch(r"-?\d+", text):
        raise ValidationError(f"Dezena nao numerica: {value}.")
    return int(text)


def _as_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        raise ValidationError("Data do sorteio vazia.")
    for pattern in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise ValidationError(f"Data do sorteio invalida: {text}.")


@dataclass(frozen=True, slots=True)
class SheetData:
    """Cabecalho normalizado e linhas. O cabecalho e o que identifica a modalidade."""

    header: tuple[str, ...]
    rows: list[dict[str, Any]]


def _rows_from_xlsx(content: bytes) -> SheetData:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    if sheet is None:
        raise ValidationError("Planilha sem aba de dados.")

    # As planilhas oficiais da Caixa declaram a dimensao da aba como uma unica celula.
    # No modo read_only o openpyxl confia nessa declaracao e devolve so o cabecalho, o que
    # faria um arquivo com milhares de concursos parecer vazio. Recalcular a dimensao
    # forca a leitura de todas as linhas de verdade.
    reset = getattr(sheet, "reset_dimensions", None)
    if callable(reset):
        reset()

    raw_rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    return _rows_from_matrix(raw_rows)


def _rows_from_csv(content: bytes) -> SheetData:
    text = content.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return _rows_from_matrix([tuple(row) for row in reader])


def _rows_from_matrix(raw_rows: Sequence[Sequence[Any]]) -> SheetData:
    header_index = _find_header(raw_rows)
    header = [normalize_header(cell) for cell in raw_rows[header_index]]
    rows: list[dict[str, Any]] = []
    for offset, raw in enumerate(raw_rows[header_index + 1 :], start=header_index + 2):
        if all(cell is None or str(cell).strip() == "" for cell in raw):
            continue
        row: dict[str, Any] = {"__line__": offset}
        for key, cell in zip(header, raw, strict=False):
            if key:
                row.setdefault(key, cell)
        rows.append(row)
    return SheetData(header=tuple(header), rows=rows)


def _find_header(raw_rows: Sequence[Sequence[Any]]) -> int:
    for index, raw in enumerate(raw_rows[:15]):
        normalized = {normalize_header(cell) for cell in raw}
        if normalized & set(CONTEST_ALIASES):
            return index
    raise ValidationError(
        "Cabecalho nao encontrado. A planilha precisa de uma coluna 'Concurso'."
    )


def _count_value_columns(header: Sequence[str], prefixes: Sequence[str]) -> tuple[str, int] | None:
    """Quantas colunas de dezenas o cabecalho traz, e sob qual prefixo.

    Conta ate a primeira ausencia: `bola1..bola7` devolve 7, nao "pelo menos 6".
    """
    for prefix in prefixes:
        present = set(header)
        total = 0
        while f"{prefix}{total + 1}" in present:
            total += 1
        if total:
            return prefix, total
    return None


def _modality_by_shape(prefix: str, count: int) -> Modality | None:
    """Que modalidade publica a planilha com esse prefixo e essa quantidade de colunas.

    Os quatro formatos com prefixo "bola" tem contagens distintas (5, 6, 7 e 15) e o
    Super Sete usa "coluna", entao o par prefixo/contagem identifica a planilha sozinho.
    """
    for candidate, strategy in STRATEGIES.items():
        if prefix in strategy.value_prefixes and strategy.value_count == count:
            return candidate
    return None


def _modality_from_filename(filename: str) -> Modality | None:
    """A modalidade que o nome do arquivo anuncia, se anunciar alguma.

    Os arquivos saem da Caixa como "Mega-Sena.xlsx" ou "Dia de Sorte.xlsx". Um arquivo
    renomeado nao acusa nada e simplesmente nao participa desta checagem.
    """
    normalized = normalize_header(re.sub(r"\.[a-z0-9]+$", "", filename, flags=re.IGNORECASE))
    # Do mais longo para o mais curto: "megasena" antes de "mega".
    for candidate in sorted(Modality, key=lambda item: -len(normalize_header(item.label))):
        if normalize_header(candidate.label) in normalized:
            return candidate
    return None


def _ensure_sheet_matches(modality: Modality, header: Sequence[str], filename: str) -> None:
    """Trava de modalidade: recusa a planilha de uma modalidade importada em outra.

    Sem isto, uma planilha do Dia de Sorte (bola1..bola7) importada na Mega-Sena passava:
    as seis primeiras colunas existem, o parser truncava a setima e as dezenas de 1 a 31
    cabem no universo de 1 a 60, entao nenhuma validacao de linha reclamava.
    """
    strategy = STRATEGIES[modality]
    found = _count_value_columns(header, strategy.value_prefixes)

    if found is not None:
        prefix, count = found
        if count != strategy.value_count:
            culpada = _modality_by_shape(prefix, count)
            detalhe = (
                f" Esse e o formato de {culpada.label}: abra a modalidade {culpada.label} "
                "para importar este arquivo."
                if culpada is not None and culpada is not modality
                else ""
            )
            raise ImportValidationError(
                f"A planilha traz {count} colunas de dezenas ({prefix}1 a {prefix}{count}) e "
                f"{modality.label} usa {strategy.value_count}.{detalhe} "
                "Nada foi alterado no banco.",
                errors=[],
            )

    anunciada = _modality_from_filename(filename)
    if anunciada is not None and anunciada is not modality:
        raise ImportValidationError(
            f"O arquivo se chama '{filename}', que e a planilha de {anunciada.label}, mas a "
            f"importacao foi aberta em {modality.label}. Abra a modalidade {anunciada.label} "
            "ou renomeie o arquivo se o nome estiver errado. Nada foi alterado no banco.",
            errors=[],
        )


def parse_file(modality: Modality, filename: str, content: bytes) -> ParseResult:
    """Le e valida o arquivo inteiro em memoria. Nao toca no banco."""
    rule = rule_for(modality)
    strategy = STRATEGIES[modality]

    lowered = filename.lower()
    if lowered.endswith(".csv") or lowered.endswith(".txt"):
        sheet = _rows_from_csv(content)
    elif lowered.endswith(".xlsx") or lowered.endswith(".xlsm"):
        sheet = _rows_from_xlsx(content)
    else:
        raise ValidationError("Formato nao suportado. Envie um arquivo .xlsx ou .csv.")

    # Antes de ler uma linha sequer: a planilha e mesmo desta modalidade?
    _ensure_sheet_matches(modality, sheet.header, filename)

    result = ParseResult(file_hash=file_hash(content))
    seen: dict[int, int] = {}

    for row in sheet.rows:
        line = int(row["__line__"])
        try:
            contest_raw = _first_present(row, CONTEST_ALIASES)
            if contest_raw is None:
                raise ValidationError("Concurso ausente.")
            contest_no = _as_int(contest_raw)
            if contest_no <= 0:
                raise ValidationError(f"Concurso invalido: {contest_no}.")
        except ValidationError as exc:
            result.errors.append(RowError(line, "concurso", str(exc)))
            continue

        if contest_no in seen:
            result.errors.append(
                RowError(
                    line,
                    "concurso",
                    f"Concurso {contest_no} duplicado (ja aparece na linha {seen[contest_no]}).",
                    str(contest_no),
                )
            )
            continue
        seen[contest_no] = line

        try:
            drawn_at = _as_date(_first_present(row, DATE_ALIASES))
        except ValidationError as exc:
            result.errors.append(RowError(line, "data", str(exc)))
            continue

        try:
            numbers = _collect_values(row, strategy.value_prefixes, strategy.value_count)
        except ValidationError as exc:
            result.errors.append(RowError(line, "dezenas", str(exc)))
            continue

        row_errors = _validate_numbers(line, numbers, rule, strategy)
        if row_errors:
            result.errors.extend(row_errors)
            continue

        extras: dict[str, Any] = {}
        if strategy.extras_reader is not None:
            try:
                extras = strategy.extras_reader(row, rule, numbers)
            except ValidationError as exc:
                result.errors.append(RowError(line, "extras", str(exc)))
                continue

        stored = tuple(numbers) if strategy.allow_repeated_values else tuple(sorted(numbers))
        result.draws.append(
            Draw(
                modality=modality,
                contest_no=contest_no,
                drawn_at=drawn_at,
                numbers=stored,
                extras=extras,
            )
        )

    result.draws.sort(key=lambda draw: draw.contest_no)
    return result


def _validate_numbers(
    line: int, numbers: Sequence[int], rule: ModalityRule, strategy: ModalityStrategy
) -> list[RowError]:
    errors: list[RowError] = []
    if len(numbers) != strategy.value_count:
        errors.append(
            RowError(
                line,
                "dezenas",
                f"Esperado {strategy.value_count} dezenas, encontrado {len(numbers)}.",
            )
        )
        return errors

    for number in numbers:
        if not rule.contains(number):
            errors.append(
                RowError(
                    line,
                    "dezenas",
                    f"Dezena {number} fora do universo "
                    f"{rule.universe_min} a {rule.universe_max}.",
                    str(number),
                )
            )

    if not strategy.allow_repeated_values and len(set(numbers)) != len(numbers):
        errors.append(RowError(line, "dezenas", "Dezenas repetidas no mesmo concurso."))

    return errors


def ensure_valid(result: ParseResult, limit: int = 50) -> None:
    """Converte a lista de erros em uma excecao de dominio com detalhe por linha."""
    if result.is_valid:
        if not result.draws:
            raise ImportValidationError(
                "O arquivo nao contem nenhum concurso valido.", errors=[]
            )
        return
    raise ImportValidationError(
        f"O arquivo tem {len(result.errors)} erro(s) de validacao. Nada foi alterado no banco.",
        errors=[error.as_dict() for error in result.errors[:limit]],
    )
